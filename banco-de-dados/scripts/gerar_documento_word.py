#!/usr/bin/env python3
"""Gera o documento Word da base de dados (docs/Documentacao-Banco-de-Dados.docx).

Junta, em um só arquivo: resumo da entrega, escolha do banco (PostgreSQL), modelagem, os DER (imagens), normalização,
demonstração de CRUD e o dicionário de dados. O texto vem dos arquivos do próprio repositório (modelagem.md,
dicionario-de-dados.md), então o Word nunca fica diferente do que está no GitHub.

Uso:   python scripts/gerar_documento_word.py
Requer: python-docx e pillow   (pip install python-docx pillow)
"""
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "docs" / "Documentacao-Banco-de-Dados.docx"
NAVY = RGBColor(0, 0, 0)  # documento todo em preto (sem cor), como no TCC anterior


# =====================================================================================================
# 1) TEXTO (Markdown simples montado a partir dos arquivos do repositório)
# =====================================================================================================
def secao(md: str, titulo_inicio: str) -> str:
    """Corpo (sem o título) da seção '## <titulo_inicio>...' de um arquivo markdown."""
    partes = re.split(r"(?m)^## ", md)
    for p in partes:
        if p.startswith(titulo_inicio):
            return "\n".join(p.split("\n")[1:]).strip()
    raise KeyError(titulo_inicio)


def rebaixar(md: str, niveis: int) -> str:
    return re.sub(r"(?m)^(#{2,4}) ", lambda m: "#" * (len(m.group(1)) + niveis) + " ", md)


def trecho_crud() -> str:
    """Trecho real da saída de crud_demo.py (bloco de fábrica: criar, ler e alterar)."""
    env = dict(os.environ, PYTHONIOENCODING="utf-8")
    saida = subprocess.run([sys.executable, str(RAIZ / "scripts" / "crud_demo.py")], capture_output=True, text=True, encoding="utf-8", env=env, check=True).stdout
    linhas = saida.splitlines()
    ini = next(i for i, l in enumerate(linhas) if l.startswith("[READ] a fábrica com suas máquinas"))
    fim = next(i for i, l in enumerate(linhas) if l.startswith("[READ] conferindo") and i > ini)
    trecho = [l if len(l) <= 118 else l[:115] + "..." for l in linhas[ini:fim + 6]]
    return "\n".join(trecho)


def montar_markdown() -> str:
    modelagem = (RAIZ / "docs" / "modelagem.md").read_text(encoding="utf-8")
    dic = (RAIZ / "docs" / "dicionario-de-dados.md").read_text(encoding="utf-8")
    dic_corpo = dic.split("\n## Módulo", 1)
    dic_modulos = "## Módulo" + dic_corpo[1]
    pre_dic = dic.split("\n## Módulo", 1)[0]
    tabela_modulos = pre_dic[pre_dic.index("| Módulo |"):].strip()

    md = f"""# Documentação da base de dados
## 1. Resumo da entrega
O Sprint 1 pede a **base de dados** do sistema: o diagrama entidade-relacionamento (DER), a própria base com as tabelas criadas e a capacidade de executar comandos de CRUD (Create, Read, Update e Delete). Tudo está no repositório do projeto, na pasta `banco-de-dados/`.

| Pedido | Onde está |
|---|---|
| Diagrama entidade-relacionamento (DER) | Pasta `der/`: visão geral, um diagrama por módulo e o modelo completo (neste documento, na seção 5) |
| A base de dados com as tabelas criadas | `digital_twin.db` (SQLite, cópia portátil) e os scripts em `sql/` (SQLite e PostgreSQL). 21 tabelas, 4 views e 345 linhas iniciais |
| Comandos de CRUD | `sql/04_crud_exemplos.sql` e `scripts/crud_demo.py`, que executa Create, Read, Update e Delete em todas as tabelas |
| Testes | `testes/test_banco.py`: 23 testes automáticos no SQLite (criação, dados, CRUD e regras de integridade) e `scripts/testar_postgres.py`: a mesma verificação num servidor PostgreSQL 16 (Docker), que passou |
| Documentação | Este documento, `docs/modelagem.md` e `docs/dicionario-de-dados.md` |

O banco foi pensado para o sistema Digital Twin Industrial, um simulador de fábrica de bebidas no estilo Cisco Packet Tracer. Ele guarda o catálogo de máquinas (com limites, capacidade, consumo, preços estimados e fontes), as fábricas montadas pelos usuários, os resultados de simulação, os planos de melhoria, o registro de eventos e a conversa com o assistente.

## 2. Banco escolhido: PostgreSQL
{rebaixar(secao(modelagem, "5. Banco escolhido"), 1).replace("O professor deixou a escolha livre. ", "A escolha do banco foi livre. ")}
## 3. O que o sistema precisa guardar
{secao(modelagem, "1. O que o sistema precisa guardar")}
## 4. Entidades e relacionamentos
{secao(modelagem, "2. Entidades e relacionamentos")}

O banco tem 21 tabelas, organizadas em 5 módulos:

{tabela_modulos}

## 5. Diagramas entidade-relacionamento (DER)
Os diagramas usam a notação "pé de galinha": `||` significa exatamente um, `|o` zero ou um e `o{{` zero ou muitos. PK é a chave primária e FK é a chave estrangeira. A visão geral mostra as 21 tabelas com as chaves; os diagramas por módulo mostram todas as colunas e são os mais fáceis de ler. As imagens também estão em alta resolução na pasta `der/` do repositório.

[[paisagem]]
### 5.1 Visão geral
![Figura 1](../der/DER-visao-geral.png)

Figura 1 – DER: visão geral das 21 tabelas, com as chaves primárias e estrangeiras.

[[retrato]]
### 5.2 Módulo Acesso
![Figura 2](../der/DER-modulo-1-acesso.png)

Figura 2 – DER do módulo Acesso (usuários).

### 5.3 Módulo Catálogo de máquinas
![Figura 3](../der/DER-modulo-2-catalogo.png)

Figura 3 – DER do módulo Catálogo de máquinas (categorias, etapas, máquinas, manuais, fontes, preços e modelos reais).

### 5.4 Módulo Apoio ao usuário
![Figura 4](../der/DER-modulo-3-apoio.png)

Figura 4 – DER do módulo Apoio ao usuário (glossário e conversa com o assistente).

### 5.5 Módulo Fábricas
![Figura 5](../der/DER-modulo-4-fabricas.png)

Figura 5 – DER do módulo Fábricas (fábricas, máquinas colocadas e conexões).

### 5.6 Módulo Simulação e resultados
![Figura 6](../der/DER-modulo-5-simulacao.png)

Figura 6 – DER do módulo Simulação e resultados (cenários, simulações, planos, passos e registro).

## 6. Normalização e decisões de projeto
{rebaixar(secao(modelagem, "3. Normalização"), 0)}
{secao(modelagem, "4. Decisões de projeto")}

## 7. CRUD: Create, Read, Update e Delete
O CRUD foi demonstrado em **todas as 21 tabelas** pelo arquivo `sql/04_crud_exemplos.sql`, que cria dados de exemplo, lê, altera e apaga, deixando o banco como estava (a conferência final conta as linhas de cada tabela e compara com o estado inicial). O script `scripts/crud_demo.py` executa cada comando e mostra o resultado, e os testes automáticos repetem o ciclo em cada tabela.

| Operação | Comando SQL | Exemplo no banco |
|---|---|---|
| Create | `INSERT` | cria uma fábrica, duas máquinas e a conexão entre elas |
| Read | `SELECT` (com `JOIN` e views) | lê a fábrica com suas máquinas e a ligação |
| Update | `UPDATE` | muda temperatura, tarifa, ritmo e unidades em paralelo |
| Delete | `DELETE` | apaga a fábrica e o que depende dela cai em cascata |

Trecho real da execução de `python scripts/crud_demo.py`:

```
{trecho_crud()}
```

O banco também recusa dados inválidos, o que é parte do CRUD seguro: uma conexão entre máquinas de fábricas diferentes, uma temperatura fora de 0 a 45 °C, uma máquina do catálogo apagada enquanto está em uso ou um nome de usuário repetido geram erro de integridade (cobertos pelos testes da classe `Integridade`).

## 8. Dicionário de dados
Cada tabela, coluna, tipo, chave e regra de integridade. O mesmo conteúdo está em `docs/dicionario-de-dados.md`, gerado automaticamente do modelo (`scripts/modelo.py`).

{rebaixar(dic_modulos, 1)}
## 9. Como executar
- **Ver o banco:** abrir `digital_twin.db` no DB Browser for SQLite (gratuito, https://sqlitebrowser.org).
- **Ver o CRUD:** `python scripts/crud_demo.py` (só precisa do Python).
- **Rodar os testes:** `python -m unittest discover -s testes -v`.
- **Criar do zero em SQLite:** `sqlite3 meu_banco.db < sql/01_schema_sqlite.sql`, depois `02_dados_iniciais.sql` e `03_views.sql`.
- **Criar em PostgreSQL com Docker:** `docker compose up -d` (sobe o servidor e já cria tudo) e `python scripts/testar_postgres.py` (confere as tabelas, roda o CRUD completo e os testes de integridade).
- **Criar num PostgreSQL já instalado:** `createdb digital_twin` e `psql digital_twin -f sql/01_schema_postgres.sql`, depois `02_dados_iniciais.sql`, `02b_ajustar_sequencias_postgres.sql` e `03_views.sql`.
- **Regerar os arquivos a partir do modelo:** `python scripts/gerar_banco.py` (SQL, banco e dicionário) e `python scripts/gerar_der.py` (diagramas).
"""
    return md


# =====================================================================================================
# 2) CONVERSÃO PARA WORD
# =====================================================================================================
doc = Document()
INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*\s][^*]*?\*)")
LARGURAS = {"Coluna": [3.0, 2.3, 3.8, 1.6, 1.6, 3.7], "Pedido": [5.0, 11.0], "Critério": [4.2, 4.2, 4.0, 3.6], "Operação": [2.2, 5.4, 8.4]}


def configurar_estilos() -> None:
    st = doc.styles
    st["Normal"].font.name, st["Normal"].font.size = "Arial", Pt(11)
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.2
    for nome, tam in (("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        h = st[nome]
        h.font.name, h.font.size, h.font.bold, h.font.color.rgb = "Arial", Pt(tam), True, NAVY
        for a in ("w:eastAsia", "w:ascii", "w:hAnsi"):
            h.element.rPr.rFonts.set(qn(a), "Arial")
        h.paragraph_format.space_before, h.paragraph_format.space_after, h.paragraph_format.keep_with_next = Pt(14), Pt(8), True
    st["List Bullet"].font.name, st["List Bullet"].font.size = "Arial", Pt(11)


def runs(par, texto, tam=None, negrito=False, italico=False) -> None:
    for tok in INLINE.split(texto):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2]); r.bold = True; r.italic = italico
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((tam or 11) - 1); r.bold = negrito
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")
            continue
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True; r.bold = negrito
        else:
            r = par.add_run(tok); r.bold = negrito; r.italic = italico
        if tam:
            r.font.size = Pt(tam)


def sombrear_paragrafo(p, cor: str) -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), cor)
    pPr.insert_element_before(shd, "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
                              "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                              "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
                              "w:rPr", "w:sectPr", "w:pPrChange")


def centralizado(texto, tam, negrito=False, cor=None, depois=6, italico=False) -> None:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(depois)
    r = p.add_run(texto); r.font.size, r.bold, r.italic = Pt(tam), negrito, italico
    if cor:
        r.font.color.rgb = cor


def campo(par, instrucao: str, texto: str) -> None:
    for tipo, conteudo in (("begin", None), ("instr", instrucao), ("separate", None), ("texto", texto), ("end", None)):
        r = par.add_run()
        if tipo == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = conteudo; r._r.append(it)
        elif tipo == "texto":
            r.text = conteudo
        else:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); r._r.append(f)


def nova_secao(paisagem: bool) -> None:
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE if paisagem else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (Cm(29.7), Cm(21)) if paisagem else (Cm(21), Cm(29.7))
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = (Cm(2), Cm(2), Cm(2), Cm(2)) if paisagem else (Cm(3), Cm(2), Cm(3), Cm(2))


def tabela(linhas: list[list[str]]) -> None:
    cab, corpo = linhas[0], linhas[1:]
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    larg = LARGURAS.get(cab[0].strip(), [16.0 / len(cab)] * len(cab))
    if len(larg) != len(cab):
        larg = [16.0 / len(cab)] * len(cab)
    for j, h in enumerate(cab):
        c = t.rows[0].cells[j]; c.text = ""
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "E7E7E7"); c._tc.get_or_add_tcPr().append(shd)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(2); runs(p, h.strip(), 9, True)
    for linha in corpo:
        cels = t.add_row().cells
        for j in range(len(cab)):
            cels[j].text = ""; p = cels[j].paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
            runs(p, (linha[j] if j < len(linha) else "").strip(), 9)
    for linha in t.rows:
        for j, c in enumerate(linha.cells):
            c.width = Cm(larg[j])
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); t.rows[0]._tr.get_or_add_trPr().append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def converter(md: str) -> None:
    linhas = md.split("\n"); i = 0; paisagem = False
    while i < len(linhas):
        s = linhas[i].strip()
        if not s or s == "---":
            i += 1; continue
        if s == "[[paisagem]]":
            nova_secao(True); paisagem = True; i += 1; continue
        if s == "[[retrato]]":
            nova_secao(False); paisagem = False; i += 1; continue
        if s.startswith("```"):
            i += 1; bloco = []
            while i < len(linhas) and not linhas[i].strip().startswith("```"):
                bloco.append(linhas[i]); i += 1
            i += 1
            for l in bloco:
                p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0; p.paragraph_format.left_indent = Cm(0.2)
                r = p.add_run(l if l else " "); r.font.name, r.font.size = "Consolas", Pt(7.5)
                r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas"); sombrear_paragrafo(p, "F5F7F9")
            doc.add_paragraph().paragraph_format.space_after = Pt(2); continue
        if s.startswith("|"):
            linhas_t = []
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                cels = linhas[i].strip().strip("|").split("|")
                if not all(re.fullmatch(r"\s*:?-{2,}:?\s*", c) for c in cels):
                    linhas_t.append(cels)
                i += 1
            tabela(linhas_t); continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
        if m:
            caminho = (RAIZ / "docs" / m.group(2)).resolve()
            w, h = Image.open(caminho).size
            largura = min(25.5, 12.8 * w / h) if paisagem else min(16.0, 21.5 * w / h)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(caminho), width=Cm(largura)); i += 1; continue
        if re.match(r"^Figura \d+ [–-] ", s):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; runs(p, s, 10, italico=True); i += 1; continue
        hm = re.match(r"^(#{1,5}) (.*)$", s)
        if hm:
            nivel = len(hm.group(1)); texto = hm.group(2).strip()
            if nivel == 1:
                i += 1; continue  # o título do documento está na capa
            p = doc.add_heading(texto.replace("`", ""), level=min(nivel - 1, 3))
            if nivel == 2 and not paisagem:
                p.paragraph_format.page_break_before = True
            i += 1; continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); runs(p, s[2:]); i += 1; continue
        mnum = re.match(r"^(\d+)\. (.*)$", s)
        if mnum:
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.6)
            runs(p, f"{mnum.group(1)}. {mnum.group(2)}"); i += 1; continue
        buf = [s]; i += 1
        while i < len(linhas) and linhas[i].strip() and not re.match(r"^(#{1,5} |\||!\[|- |\d+\. |```|---|\[\[|Figura \d+ [–-] )", linhas[i].strip()):
            buf.append(linhas[i].strip()); i += 1
        p = doc.add_paragraph(); texto = " ".join(buf)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(texto) > 140 else WD_ALIGN_PARAGRAPH.LEFT
        runs(p, texto)


def main() -> None:
    configurar_estilos()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2), Cm(3), Cm(2)
    # Capa no padrão do TCC anterior: Arial 12 negrito, preto, centralizada, em maiúsculas (sem número de página)
    sec.different_first_page_header_footer = True
    def capa_linha(texto, antes=0, depois=0):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before, p.paragraph_format.space_after, p.paragraph_format.line_spacing = Pt(antes), Pt(depois), 1.5
        r = p.add_run(texto); r.font.name, r.font.size, r.bold = "Arial", Pt(12), True
        r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial")
    capa_linha("FACULDADE DE TECNOLOGIA DE COTIA")
    capa_linha("ENZO FERNANDES", 109)
    capa_linha("VICTOR RAFAEL DE OLIVEIRA DIAS")
    capa_linha("DIGITAL TWIN INDUSTRIAL", 173)
    capa_linha("BASE DE DADOS: DER, MODELAGEM E CRUD")
    capa_linha("Cotia - SP", 239)
    capa_linha("2026")
    p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
    r = p.add_run("Sumário"); r.bold, r.font.size, r.font.color.rgb = True, Pt(15), NAVY
    campo(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', "Atualize o campo para ver o sumário (clique com o botão direito > Atualizar campo).")
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    campo(fp, "PAGE", "1")
    converter(montar_markdown())
    cfg = doc.settings.element
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
    cfg.insert_element_before(uf, "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids", "m:mathPr", "w:attachedSchema",
                              "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade",
                              "w:captions", "w:readModeInkLockDown", "w:smartTagType", "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")
    for z in cfg.findall(qn("w:zoom")):
        if z.get(qn("w:percent")) is None:
            z.set(qn("w:percent"), "100")
    doc.core_properties.title = "Digital Twin Industrial: documentação da base de dados"
    doc.core_properties.author = "Victor Rafael de Oliveira Dias; Enzo Fernandes"
    doc.save(SAIDA)
    print(f"salvo {SAIDA.relative_to(RAIZ)} ({SAIDA.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
